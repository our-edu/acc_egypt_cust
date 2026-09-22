import io

import frappe
from frappe import _


PARTY_ACCOUNT_TYPES = ("Receivable", "Payable")


def _get_employee_loan_accounts(company):
    if not company or not frappe.db.exists("DocType", "OurEdu HR Setting"):
        return frozenset()

    loan_account = frappe.db.get_value(
        "OurEdu HR Setting", {"company": company}, "employee_loan_account"
    )
    return frozenset([loan_account]) if loan_account else frozenset()


def _is_employee_party(row):
    return (row.custom_party_type or row.party_type) == "Employee"


def _should_sync_party_for_row(row, account_type, loan_accounts):
    if _is_employee_party(row):
        return True
    if account_type in PARTY_ACCOUNT_TYPES:
        return True
    return row.account in loan_accounts


def set_title_to_name(doc, method=None):
    """Set Journal Entry title to its document name (ID)."""
    doc.db_set("title", doc.name)


def auto_submit_depreciation_entry(doc, method=None):
    """Force-submit Asset Depreciation Journal Entries even if a Journal Entry
    workflow would otherwise leave them as drafts (erpnext skips its own
    auto-submit for these once a workflow is configured on Journal Entry)."""
    if doc.voucher_type == "Depreciation Entry" and doc.docstatus == 0 and doc.meta.get_workflow():
        doc.flags.ignore_permissions = True
        doc.submit()


def _resolve_party_name(party_type, party):
    """Fetch the display name for a party, mirroring erpnext's Payment Entry convention."""
    if not party_type or not party:
        return None

    fieldname = "title" if party_type == "Shareholder" else party_type.lower() + "_name"
    if frappe.db.has_column(party_type, fieldname):
        return frappe.db.get_value(party_type, party, fieldname)
    return frappe.db.get_value(party_type, party, "name")


@frappe.whitelist()
def get_party_name(party_type, party):
    return _resolve_party_name(party_type, party)


def sync_custom_party_to_party(doc, method=None):
    """Copy custom_party fields to party_type/party for AR/AP, loan, and Employee party rows."""
    loan_accounts = _get_employee_loan_accounts(doc.company)

    for row in doc.get("accounts"):
        if not row.account:
            row.party_type = None
            row.party = None
            row.custom_party_name = None
            continue

        account_type = frappe.get_cached_value("Account", row.account, "account_type")

        if _should_sync_party_for_row(row, account_type, loan_accounts):
            row.party_type = row.custom_party_type or row.party_type or None
            row.party = row.custom_party or row.party or None
        else:
            row.party_type = None
            row.party = None

        row.custom_party_name = _resolve_party_name(row.party_type, row.party)


@frappe.whitelist()
def parse_je_accounts_excel(file_url, company=None):
    """
    Parse an uploaded Excel file and return rows for Journal Entry accounts table.
    Expected columns (case-insensitive): account, debit, credit, party_type, party,
    cost_center, project, user_remark, reference_no, reference_date, multi_currency,
    currency, exchange_rate, debit_in_account_currency, credit_in_account_currency.

    Returns {"rows": [...], "multi_currency": bool}.

    "multi_currency" is a Journal Entry (parent) checkbox, not a Journal Entry
    Account (child) field, so it can't be assigned into a row dict like the rest of
    FIELD_MAP. It's read per row here and folded into a single overall flag: true if
    any row's cell is truthy, OR if any row's account is booked in a currency other
    than the company's (via `company`, passed in from frm.doc.company) -- this way an
    imported foreign-currency line doesn't need the sheet author to remember to tick
    the column, and doesn't hit erpnext's "Please check Multi Currency option to
    allow accounts with other currency" throw on save.
    """
    try:
        import openpyxl
    except ImportError:
        frappe.throw(_("openpyxl is required. Run: bench pip install openpyxl"))

    # Resolve the file from the Frappe File doctype
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    file_path = file_doc.get_full_path()

    with open(file_path, "rb") as f:
        content = f.read()

    wb = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
    ws = wb.active

    # Read header row (first row)
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        val = cell.value
        headers.append(str(val).strip().lower() if val is not None else "")

    # Mapping from Excel header -> JE Account field. "multi_currency" is handled
    # separately below since it lives on the parent Journal Entry, not this child row.
    FIELD_MAP = {
        "account": "account",
        "debit": "debit_in_account_currency",
        "debit_in_account_currency": "debit_in_account_currency",
        "credit": "credit_in_account_currency",
        "credit_in_account_currency": "credit_in_account_currency",
        "exchange_rate": "exchange_rate",
        "currency": "account_currency",
        "account_currency": "account_currency",
        "party_type": "custom_party_type",
        "party": "custom_party",
        "custom_party_type": "custom_party_type",
        "custom_party": "custom_party",
        "cost_center": "cost_center",
        "project": "project",
        "user_remark": "user_remark",
        "remark": "user_remark",
        "reference_type": "reference_type",
        "reference_name": "reference_name",
        "reference_no": "custom_reference_no",
        "custom_reference_no": "custom_reference_no",
        "reference_date": "custom_reference_date",
        "custom_reference_date": "custom_reference_date",
    }
    MULTI_CURRENCY_HEADER = "multi_currency"

    company_currency = (
        frappe.get_cached_value("Company", company, "default_currency") if company else None
    )

    def _is_truthy_cell(val):
        if isinstance(val, bool):
            return val
        if val is None:
            return False
        return str(val).strip().lower() in ("1", "true", "yes", "y", "x")

    rows = []
    multi_currency = False
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Skip entirely blank rows
        if all(v is None or str(v).strip() == "" for v in row):
            continue

        entry = {}
        row_multi_currency = False
        for idx, header in enumerate(headers):
            if idx >= len(row):
                break
            val = row[idx]

            if header == MULTI_CURRENCY_HEADER:
                if _is_truthy_cell(val):
                    row_multi_currency = True
                continue

            field = FIELD_MAP.get(header)
            if not field or val is None:
                continue
            # Numeric fields: coerce to float
            if field in ("debit_in_account_currency", "credit_in_account_currency"):
                try:
                    val = float(val)
                except (ValueError, TypeError):
                    val = 0.0
            elif field == "exchange_rate":
                # Leave unset on a bad/blank cell so erpnext auto-fetches it on save
                # (Journal Entry.set_exchange_rate) instead of forcing a wrong rate.
                try:
                    val = float(val)
                except (ValueError, TypeError):
                    continue
            # Date fields: openpyxl returns datetime objects for date cells
            elif field == "custom_reference_date":
                import datetime
                if isinstance(val, (datetime.datetime, datetime.date)):
                    val = val.strftime("%Y-%m-%d")
                else:
                    val = str(val).strip()
            else:
                val = str(val).strip()
            entry[field] = val

        if not entry.get("account"):
            continue

        # Auto-detect a foreign-currency account regardless of the sheet's own
        # multi_currency column (see docstring above).
        if company_currency:
            account_currency = frappe.get_cached_value("Account", entry["account"], "account_currency")
            if account_currency and account_currency != company_currency:
                row_multi_currency = True

        if row_multi_currency:
            multi_currency = True

        rows.append(entry)

    if not rows:
        frappe.throw(_("No valid rows found. Make sure the file has an 'account' column with data."))

    return {"rows": rows, "multi_currency": multi_currency}


@frappe.whitelist(allow_guest=False)
def download_je_accounts_template():
    """Return an xlsx template file for Journal Entry accounts import."""
    try:
        import openpyxl
    except ImportError:
        frappe.throw(_("openpyxl is required. Run: bench pip install openpyxl"))

    columns = [
        "account",
        "debit",
        "credit",
        "party_type",
        "party",
        "cost_center",
        "project",
        "user_remark",
        "reference_no",
        "reference_date",
        "multi_currency",
        "currency",
        "exchange_rate",
        "debit_in_account_currency",
        "credit_in_account_currency",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Journal Entry Accounts"
    ws.append(columns)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    frappe.local.response.filename = "je_accounts_template.xlsx"
    frappe.local.response.filecontent = buf.read()
    frappe.local.response.type = "download"
    frappe.local.response.content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )



